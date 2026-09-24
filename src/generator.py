"""Gerador de gabarito de escala de leitores para um determinado mês."""

import calendar
import datetime

import openpyxl
from openpyxl.styles import (  # pyright: ignore[reportMissingTypeStubs]
    Alignment,
    Font,
    PatternFill,
)

DIAS_SEMANA_MAP = {
    2: ["19:00"],  # Quarta
    3: ["19:00"],  # Quinta
    4: ["19:00"],  # Sexta
    5: ["19:00"],  # Sábado
    6: ["07:30", "09:00", "19:00"],  # Domingo
}

DIAS_PT = [
    "Segunda-feira",
    "Terça-feira",
    "Quarta-feira",
    "Quinta-feira",
    "Sexta-feira",
    "Sábado",
    "Domingo",
]


def _formatar_linha(ws, row_idx: int, valores: list[str], fill_color: str) -> None:
    for col_idx, valor in enumerate(valores, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=valor)
        cell.font = Font(name="Calibri", size=10)
        cell.fill = PatternFill(
            start_color=fill_color, end_color=fill_color, fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")


def gerar_gabarito_mes(ano: int, mes: int, output_path: str):
    """Gera e salva uma planilha com as celebrações do mês informado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Escala {mes:02d}-{ano}"

    # Título do mês
    nome_mes = calendar.month_name[mes].upper()
    ws.merge_cells("A1:H1")
    ws["A1"] = f"ESCALA DE LEITORES - {nome_mes} DE {ano} (GABARITO)"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(
        start_color="1A365D", end_color="1A365D", fill_type="solid"
    )
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "DATA",
        "DIA",
        "HORA",
        "COMENTARISTA",
        "1ª LEITURA",
        "SALMO",
        "2ª LEITURA",
        "PRECES",
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="2B6CB0", end_color="2B6CB0", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")

    num_dias = calendar.monthrange(ano, mes)[1]
    row_idx = 3

    for dia in range(1, num_dias + 1):
        dt = datetime.date(ano, mes, dia)
        weekday = dt.weekday()

        if weekday in DIAS_SEMANA_MAP:
            for hora in DIAS_SEMANA_MAP[weekday]:
                is_weekend = weekday in [5, 6]
                fill_color = "EDF2F7" if is_weekend else "FFFFFF"

                row_values = [
                    dt.strftime("%d/%m/%Y"),
                    DIAS_PT[weekday],
                    hora,
                    "",
                    "",
                    "",
                    "",
                    "",  # Campos em branco
                ]

                _formatar_linha(ws, row_idx, row_values, fill_color)
                row_idx += 1

    wb.save(output_path)
